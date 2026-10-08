from io import BytesIO
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from thermo import process_path, FREE

WINTER_NOTE = (
    "Die relative Feuchte bezieht sich in dieser App immer auf flüssiges Wasser. "
    "Unter 0 °C verwenden wir dazu unterkühltes Wasser als Rechenbezug. "
    "Dadurch bleibt die Bedeutung der Prozentangabe beim Erwärmen winterlicher Außenluft gleich. "
    "Das bedeutet nicht, dass Wasser in einer realen Anlage unter 0 °C flüssig bleiben muss: "
    "Reif, Eis und die Vereisung eines Kühlers werden hier nicht berechnet. "
    "Kalte Außenluft kann eingegeben und erwärmt werden. Prozesse mit Wasserabscheidung unter 0 °C sind ausgeschlossen. "
    "Bei Mess- oder Wetterdaten muss geprüft werden, ob deren Feuchteangabe ebenfalls auf Wasser bezogen ist."
)


def excel_export(rows, states, p):
    points, processes = [], []
    for i, (row, s) in enumerate(zip(rows, states), 1):
        data = {"Punkt": i, "Name": row["name"], "Vorgabe 1": row["pair"][0], "Wert 1": row["values"][0],
                "Vorgabe 2": row["pair"][1], "Wert 2": row["values"][1], "Druck [hPa]": p}
        if s:
            data.update({"T [°C]": s.T, "x [g/kg trockene Luft]": s.x, "φ [% Wasserbezug]": s.phi,
                         "ρ [kg/m³ feuchte Luft]": s.rho, "h [kJ/kg trockene Luft]": s.h,
                         "Taupunkt Wasser [°C]": s.dew, "Status": "Gültig"})
        else:
            data["Status"] = "Ungültig bei aktuellem Druck"
        points.append(data)
        if i > 1:
            a, b = states[i-2], s
            entry = {"Von": i-1, "Nach": i, "Prozess": row["process"]}
            if a and b:
                try:
                    _, water = process_path(a, b, row["process"], p)
                    entry.update({"ΔT [K]": b.T-a.T, "Δx [g/kg trockene Luft]": b.x-a.x,
                                  "Δh Luft [kJ/kg trockene Luft]": b.h-a.h,
                                  "Abgeschiedenes Wasser [g/kg trockene Luft]": water,
                                  "Status": "Nur geometrische Verbindung" if row["process"] == FREE else "Gültig"})
                except ValueError as exc:
                    entry["Status"] = str(exc)
            else:
                entry["Status"] = "Ungültiger Endpunkt"
            processes.append(entry)
    notes = pd.DataFrame([
        ("Winterkonvention", WINTER_NOTE),
        ("Einheiten", "x und h beziehen sich auf 1 kg trockene Luft; ρ auf die gesamte feuchte Luft pro m³."),
        ("Druck [hPa]", str(p)),
        ("Diagramm", "Fest: −15 bis 40 °C; 0 bis 20 g/kg trockene Luft."),
        ("Kühlmodell", "Ideale Abkühlung: x konstant bis Taupunkt, danach gesättigte Luft mit kontinuierlicher Kondensatabscheidung. Kein reales Kühler-/Bypassmodell."),
        ("Prozessgrenzen", "Isotherme Be-/Entfeuchtung benötigt Wärmeausgleich. Isenthalpe Befeuchtung ist eine Näherung; Mischprozesse sind nicht enthalten."),
        ("Energie", "Δh ist die Enthalpieänderung der Luft. Sie ist bei Kondensatabscheidung nicht die vollständige Kühlerwärmebilanz; Kondensatenthalpie und Verluste sind nicht bilanziert."),
        ("Stoffwerte", "Ideales Gasgemisch, ASHRAE 2017/PsychroLib. Sättigungsdruck unter 0 °C: Murphy & Koop 2005, Gl. 10, stetig an ASHRAE bei 0 °C angeschlossen."),
        ("Quelle ASHRAE/PsychroLib", "https://psychrometrics.github.io/psychrolib/api_docs.html"),
        ("Quelle Murphy & Koop", "https://doi.org/10.1256/qj.04.94"),
    ], columns=["Thema", "Beschreibung"])
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        pd.DataFrame(points).to_excel(writer, sheet_name="Zustände", index=False)
        pd.DataFrame(processes).to_excel(writer, sheet_name="Prozesse", index=False)
        notes.to_excel(writer, sheet_name="Modell und Hinweise", index=False)
        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="007B83")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[1].height = 42
            for col in sheet.columns:
                sheet.column_dimensions[col[0].column_letter].width = min(42, max(15, len(str(col[0].value))+2))
                for cell in col[1:]:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    if cell.data_type == "f":
                        cell.data_type = "s"  # User-supplied names must never become Excel formulas.
                    if isinstance(cell.value, float):
                        cell.number_format = "0.000"
            if sheet.title == "Modell und Hinweise":
                sheet.column_dimensions["B"].width = 105
                for i in range(2, sheet.max_row+1):
                    sheet.row_dimensions[i].height = 75 if i == 2 else 48
    return buffer.getvalue()
